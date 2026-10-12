"""Export an author-approved Array revision without regenerating its contents.

Requires python-docx, pypdf, and an existing pdflatex installation. Sources are
copied before compilation; auxiliary files never enter the working manuscript
directory, and the annotated review PDF is never an export destination.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from pypdf import PdfReader


PAPER = Path(__file__).resolve().parents[1]
REVISION = PAPER / "array_revision"
ANALYSIS = PAPER / "array_revision_fair" / "analysis" / "full"
TITLE = "A few farther-sighted trees can improve shallow forests"
ANNOTATED_SHA256 = "1f8ded758b6970b6348424ff0bbcf50c270f28006c99f78fb60add8b3dd85723"


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def document(title: str) -> Document:
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin = section.bottom_margin = Mm(21)
    section.left_margin = section.right_margin = Mm(23)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.0
    doc.styles["Title"].font.name = "Times New Roman"
    doc.styles["Title"].font.size = Pt(15)
    doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
    # The bundled Word template contains a decorative Title border and theme font.
    for border in doc.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for style_id in ("Normal", "Title", "TitleChar"):
        for style in doc.styles:
            if style.style_id != style_id:
                continue
            style.font.name = "Times New Roman"
            for fonts in style.element.xpath(".//w:rFonts"):
                for attribute in list(fonts.attrib):
                    if "theme" in attribute.lower():
                        del fonts.attrib[attribute]
            for spacing in style.element.xpath("./w:rPr/w:spacing"):
                spacing.set(qn("w:val"), "0")
    doc.core_properties.author = "Mathias Valla"
    doc.core_properties.title = title
    doc.core_properties.subject = "Array revision ARRAY-D-26-03529"
    doc.core_properties.comments = ""
    return doc


def make_editable_files(destination: Path, highlights: list[str]) -> None:
    doc = document("Cover letter")
    text = (REVISION / "cover_letter.txt").read_text()
    for block in text.strip().split("\n\n"):
        # Preserve the intentional line breaks in addresses and the signature.
        lines = block.splitlines()
        if block.startswith(("Dr Zaved", "Yours sincerely,")):
            paragraph = doc.add_paragraph("\n".join(lines))
        else:
            paragraph = doc.add_paragraph(" ".join(lines))
        if block.startswith("Re:"):
            paragraph.runs[0].bold = True
    doc.save(destination / "cover_letter.docx")

    doc = document("Article highlights")
    doc.add_paragraph("Article highlights", "Title")
    paragraph = doc.add_paragraph(TITLE)
    paragraph.paragraph_format.space_after = Pt(16)
    for bullet in highlights:
        paragraph = doc.add_paragraph(bullet[2:], "List Bullet")
        paragraph.paragraph_format.space_after = Pt(10)
    doc.save(destination / "article_highlights.docx")
    for filename in ("cover_letter.txt", "highlights.txt", "response_to_reviewers.tex"):
        shutil.copy2(REVISION / filename, destination / filename)


def compile_pdf(source: Path, build: Path, destination: Path, pdflatex: str) -> dict:
    build.mkdir(parents=True)
    command = [pdflatex, "-interaction=nonstopmode", "-halt-on-error",
               "-file-line-error", f"-output-directory={build}", str(source)]
    environment = dict(os.environ, LC_ALL="C", LANG="C")
    for pass_number in range(1, 4):
        run = subprocess.run(command, cwd=source.parent, env=environment,
                             text=True, capture_output=True, check=False)
        (build / f"pass_{pass_number}.txt").write_text(run.stdout + run.stderr)
        if run.returncode:
            raise RuntimeError(f"LaTeX failed for {source}; diagnostics: {build}")
    log = (build / f"{source.stem}.log").read_text(errors="replace")
    errors = ("Undefined control sequence", "There were undefined references",
              "There were undefined citations", "Rerun to get cross-references right")
    if any(message in log for message in errors):
        raise RuntimeError(f"Unresolved LaTeX diagnostics: {build}")
    shutil.copy2(build / f"{source.stem}.pdf", destination)
    reader = PdfReader(destination)
    text = "\n".join(page.extract_text() or "" for page in reader.pages)
    if "??" in text or ("ARRAY-D-26-03529" not in text and source.stem != "main"):
        raise RuntimeError(f"Unexpected PDF text in {destination}")
    if any(page.get("/Annots") and any(
        item.get_object().get("/Subtype") not in ("/Link",)
        for item in page["/Annots"]
    ) for page in reader.pages):
        raise RuntimeError(f"Unexpected review annotations in {destination}")
    return {
        "pages": len(reader.pages),
        "unresolved_references": False,
        "overfull_boxes": re.findall(r"Overfull [^\n]+", log),
        "build_directory": str(build),
    }


def zip_files(destination: Path, paths: list[Path], base: Path) -> None:
    with ZipFile(destination, "w", ZIP_DEFLATED) as archive:
        for path in sorted(paths):
            archive.write(path, path.relative_to(base).as_posix())
    with ZipFile(destination) as archive:
        if archive.testzip() is not None:
            raise RuntimeError(f"Corrupt submission archive: {destination}")


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--destination", required=True, type=Path)
    full_tex = Path("/usr/local/texlive/2025/bin/universal-darwin/pdflatex")
    parser.add_argument("--pdflatex", default=str(full_tex) if full_tex.exists() else "pdflatex")
    args = parser.parse_args()
    destination = args.destination.resolve()
    if not destination.is_relative_to(REVISION.resolve()):
        raise ValueError("Use a new submission directory inside array_revision")
    annotated = REVISION / "output" / "pdf" / "revised_manuscript.pdf"
    if destination == annotated.parent or annotated.is_relative_to(destination):
        raise ValueError("The annotated review-copy directory is protected")
    if destination.exists():
        raise FileExistsError(f"Refusing to replace an existing packet: {destination}")
    if digest(annotated) != ANNOTATED_SHA256:
        raise ValueError("Annotated PDF changed; confirm the preservation baseline")

    manuscript = (REVISION / "main.tex").read_text()
    response = (REVISION / "response_to_reviewers.tex").read_text()
    if TITLE not in manuscript or TITLE not in response:
        raise ValueError("Manuscript and response titles do not match")
    citations = set()
    for group in re.findall(r"\\cite\w*\{([^}]+)\}", manuscript):
        citations.update(group.split(","))
    bibliography = re.findall(r"\\bibitem\{([^}]+)\}", manuscript)
    if citations != set(bibliography) or len(bibliography) != len(set(bibliography)):
        raise ValueError("Bibliography closure or uniqueness check failed")
    highlights = (REVISION / "highlights.txt").read_text().strip().splitlines()
    if not 3 <= len(highlights) <= 5 or any(
        not line.startswith("- ") or len(line) > 85 for line in highlights
    ):
        raise ValueError("Highlights must contain 3-5 bullets of at most 85 characters")

    validation = json.loads((ANALYSIS / "validation.json").read_text())
    if not validation["validated_complete_cohort"] or validation["missing_or_skipped_tasks"]:
        raise ValueError("Only the complete validated fair experiment may be exported")
    for name, expected in validation["output_sha256"].items():
        if digest(ANALYSIS / name) != expected:
            raise ValueError(f"Validated evidence changed: {name}")

    destination.mkdir(parents=True)
    source_dir = destination / "manuscript_source"
    source_dir.mkdir()
    for name in ("main.tex", "elsarticle.cls", "elsarticle-num.bst"):
        shutil.copy2(REVISION / name, source_dir / name)
    (source_dir / "README.txt").write_text(
        "Array revision ARRAY-D-26-03529\n"
        "A few farther-sighted trees can improve shallow forests\n\n"
        "Compile main.tex with pdfLaTeX at least twice. The three figures,\n"
        "two tables, and numerical bibliography are inline; no BibTeX pass\n"
        "or external figure is required. elsarticle.cls retains its license.\n"
        "This is a snapshot of the author-approved working source, not an\n"
        "instruction to regenerate the article from historical assemblers.\n"
    )
    make_editable_files(destination, highlights)
    pdf_dir = destination / "output" / "pdf"
    pdf_dir.mkdir(parents=True)
    temporary = Path(tempfile.mkdtemp(prefix="array_submission_", dir="/private/tmp"))
    pdf_checks = {
        "manuscript": compile_pdf(source_dir / "main.tex", temporary / "manuscript",
                                  pdf_dir / "revised_manuscript.pdf", args.pdflatex),
        "response": compile_pdf(destination / "response_to_reviewers.tex", temporary / "response",
                                pdf_dir / "response_to_reviewers.pdf", args.pdflatex),
    }
    zip_files(destination / "manuscript_source.zip", list(source_dir.iterdir()), source_dir)
    shutil.copy2(REVISION / "graphical_abstract_fair.tiff", destination / "graphical_abstract.tiff")
    shutil.copy2(REVISION / "SUBMISSION_CHECKLIST.md", destination / "README.md")
    if digest(source_dir / "main.tex") != digest(REVISION / "main.tex"):
        raise RuntimeError("Working manuscript changed during export")
    if digest(annotated) != ANNOTATED_SHA256:
        raise RuntimeError("Annotated review PDF was not preserved")
    manifest = {
        "manuscript_reference": "ARRAY-D-26-03529",
        "title": TITLE,
        "date": "2026-10-12",
        "protocol_fingerprint": validation["protocol_fingerprint"],
        "datasets": validation["datasets"],
        "outer_tasks": validation["outer_tasks"],
        "all_model_rows_total": validation["all_model_rows_total"],
        "highlight_character_counts_including_bullet": [len(line) for line in highlights],
        "bibliography_entries": len(bibliography),
        "annotated_review_pdf_preserved_sha256": ANNOTATED_SHA256,
        "pdf_checks": pdf_checks,
        "files_sha256": {
            path.relative_to(destination).as_posix(): digest(path)
            for path in sorted(destination.rglob("*")) if path.is_file()
        },
        "status": "Exported; visual review and author portal checks required before upload",
        "submission_made": False,
    }
    (destination / "submission_manifest.json").write_text(json.dumps(manifest, indent=2) + "\n")
    zip_files(destination / "array_revision_submission.zip",
              [path for path in destination.rglob("*") if path.is_file()], destination)
    print(json.dumps({"destination": str(destination), "pdf_checks": pdf_checks,
                      "highlights": manifest["highlight_character_counts_including_bullet"]}, indent=2))


if __name__ == "__main__":
    main()
